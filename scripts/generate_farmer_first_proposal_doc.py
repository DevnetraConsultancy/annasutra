"""Generates 'AnnaSutra_2_0_Farmer_First_Architecture_Proposal.docx'.

Typography rules (to avoid the box-character problem seen before):
- Arial for text, Courier New for code, ASCII-only content, no emoji.
- Diagrams are pre-rendered PNGs (scripts/arch_proposal_diagrams.py).

Run:
    backend/venv/Scripts/python.exe scripts/arch_proposal_diagrams.py
    backend/venv/Scripts/python.exe scripts/generate_farmer_first_proposal_doc.py
"""

import os
import shutil
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DIAG = os.path.join(ROOT, "docs", "diagrams", "proposal")
OUT_NAME = "AnnaSutra_2_0_Farmer_First_Architecture_Proposal.docx"
OUT_DOCS = os.path.join(ROOT, "docs", OUT_NAME)
OUT_ROOT = os.path.join(ROOT, OUT_NAME)

GREEN = RGBColor(11, 61, 46)
SAFFRON = RGBColor(194, 103, 10)
TEXT = RGBColor(30, 41, 59)
MUTED = RGBColor(100, 116, 139)

A4_W, A4_H = Inches(8.27), Inches(11.69)


# ------------------------------------------------------------------ helpers
def shade(cell, hex_color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


def margins(cell, top=70, bottom=70, left=110, right=110):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'))


def borders(table, color="CBD5E1"):
    tbl = table._tbl
    pr = tbl.tblPr
    pr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        + "".join(f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                  for e in ("top", "left", "bottom", "right", "insideH", "insideV"))
        + "</w:tblBorders>"))


def run(p, text, bold=False, italic=False, size=10, color=TEXT, font="Arial"):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = font
    r._element.rPr.rFonts.set(qn_east_asia(), font)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return r


def qn_east_asia():
    from docx.oxml.ns import qn
    return qn("w:eastAsia")


def section(doc, landscape, margin=0.7):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = A4_H, A4_W
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = A4_W, A4_H
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(margin))
    return sec


def h1(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run(p, (f"{num}  " if num else "") + text, bold=True, size=17, color=GREEN)
    # underline rule
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="C2670A"/></w:pBdr>'))
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, size=12, color=GREEN)
    return p


def para(doc, text, size=10, italic=False, color=TEXT, after=5, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        run(p, bold_lead + " ", bold=True, size=size, color=color)
    run(p, text, size=size, italic=italic, color=color)
    return p


def bullets(doc, items, size=10, bold_lead=True):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.1
        run(p, "- ", bold=True, size=size, color=SAFFRON)
        if bold_lead and "::" in it:
            lead, rest = it.split("::", 1)
            run(p, lead.strip() + ": ", bold=True, size=size)
            run(p, rest.strip(), size=size)
        else:
            run(p, it, size=size)


def table(doc, headers, rows, widths, size=8.5, head_fill="0B3D2E", first_col_bold=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        shade(c, head_fill)
        margins(c)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(c.paragraphs[0], h, bold=True, size=size, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.width = Inches(widths[i])
            margins(c)
            if ri % 2:
                shade(c, "F1F5F9")
            lines = str(val).split("\n")
            for li, ln in enumerate(lines):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                run(p, ln, size=size, bold=(first_col_bold and i == 0))
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, title, text, width=6.85, fill="F0FDF4", bar="0B3D2E"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.width = Inches(width)
    shade(c, fill)
    margins(c, 110, 110, 180, 160)
    c._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="30" w:space="0" w:color="{bar}"/>'
        f'<w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'))
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, title.upper() + "  ", bold=True, size=9, color=GREEN)
    run(p, text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(doc, text, width=6.85, size=8):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.width = Inches(width)
    shade(c, "F1F5F9")
    margins(c, 100, 100, 160, 120)
    c._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="C2670A"/><w:bottom w:val="single" w:sz="4" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:color="CBD5E1"/></w:tcBorders>'))
    lines = text.strip("\n").split("\n")
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run(p, ln if ln else " ", size=size, font="Courier New", color=RGBColor(15, 23, 42))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, fname, caption, width):
    path = os.path.join(DIAG, fname)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(c, "Figure: " + caption, italic=True, size=8.5, color=MUTED)


def footer(sec):
    f = sec.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "AnnaSutra 2.0  |  Farmer-First Architecture Proposal  |  Page ", size=8, color=MUTED)
    r = p.add_run()
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    r2 = p.add_run()
    r2._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'))
    r3 = p.add_run()
    r3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


PW = 6.85   # portrait usable width (in)
LW = 10.25  # landscape usable width (in)


# ------------------------------------------------------------------ content
def entity_chapter(doc, num, title, promise, does, never, invariants, kpis, pseudo, fig, fig_caption):
    section(doc, landscape=False)
    h1(doc, title, num)
    callout(doc, "Promise to this entity", promise)
    table(doc, ["Does", "Never does"],
          [["\n".join("- " + x for x in does), "\n".join("- " + x for x in never)]],
          [PW / 2, PW / 2], size=9)
    h2(doc, "Algorithm (pseudocode)")
    code(doc, pseudo)
    h2(doc, "Invariants the system enforces")
    bullets(doc, invariants, size=9.5, bold_lead=False)
    h2(doc, "How we know it works (KPIs)")
    bullets(doc, kpis, size=9.5, bold_lead=False)
    section(doc, landscape=True, margin=0.55)
    figure(doc, fig, fig_caption, LW)


def build():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    sec0 = doc.sections[0]
    sec0.page_width, sec0.page_height = A4_W, A4_H
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec0, side, Inches(0.7))
    footer(sec0)

    # ---------------- cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    run(p, "PROPOSAL FOR REVIEW", bold=True, size=10, color=SAFFRON)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run(p, "AnnaSutra 2.0", bold=True, size=34, color=GREEN)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run(p, "Farmer-First Trade and Traceability Architecture", bold=True, size=20, color=TEXT)
    para(doc, "Core process, entity-by-entity algorithms and block diagrams for a platform whose first customer is the farmer.",
         size=12, color=MUTED, after=18)
    table(doc, ["Item", "Detail"], [
        ["Prepared", "22 October 2026"],
        ["Basis", "Review of the current build (git commit 69b5878): lots, handover, tolerance, disputes, hash-chained ledger, 11 languages"],
        ["Audience", "Product owner, architects, FPO and mandi partners, pilot sponsors"],
        ["Status", "Design proposal. Nothing in this document has been built or validated with farmers yet."],
        ["How to read", "Sections 1-6 explain the core concept. Section 7 gives one chapter and one block diagram per entity. Sections 8-14 cover special flows, scoring, money, data, roadmap and risks."],
    ], [1.3, 5.55], size=9, first_col_bold=True)
    callout(doc, "Honest framing",
            "No design is flawless. This proposal is built to fail safely: every known weakness is listed in Section 14 with a containment, "
            "and every number is labelled as a starting hypothesis to be tuned in a pilot.", fill="FFF7E6", bar="C2670A")

    # ---------------- 1 executive summary
    section(doc, landscape=False)
    h1(doc, "Executive summary", "1")
    para(doc, "AnnaSutra today proves where a lot of produce travelled: lots, a two-step handover, tolerance checks and a tamper-evident ledger. "
              "That is valuable to consumers and regulators. For a farmer it is not yet valuable, because the system records custody and weight "
              "but not money, and it quietly treats every farmer problem as a data-entry problem.", bold_lead="Where we are.")
    para(doc, "A traceability system survives only if the person who does the most work, the farmer, is paid faster, cheated less and rewarded for proof. "
              "So the design inverts the priority: money and trust come first, and traceability becomes the by-product of honest trade.",
         bold_lead="The thesis.")
    h2(doc, "Seven structural moves")
    bullets(doc, [
        "Make the Trade the central object:: price, quantity, payment terms and delivery terms are recorded together. A handover becomes the delivery step of a trade.",
        "Separate title from custody:: carriers and warehouses hold goods but never own them, and transit loss is attributed by agreed terms, never silently to the farmer.",
        "Pay the undisputed amount immediately:: only the contested part waits. Money is never held hostage to a small disagreement.",
        "Remove the pre-link requirement:: any verified or provisional party can trade through an open handshake that works with no network.",
        "Make evidence two-sided:: both parties sign the weigh slip; a rejection needs a reason and a photo; unevidenced rejections do not stop payment.",
        "Show reputation before dispatch:: the farmer sees the buyer's Pay-Score; buyers see the farmer's trust band. Scores unlock benefits and never gate market access.",
        "Return value to the farmer:: mandi price band at listing, a consent-based credit passport, and a visible share of the consumer rupee.",
    ], size=10)
    h2(doc, "Outcome targets (hypotheses to prove in a pilot, not promises)")
    table(doc, ["Metric", "Why it matters to a farmer", "Starting target"], [
        ["Delivery-to-payment time", "Cash flow is the farmer's biggest daily constraint", "Median under 24 h; 90% of cash and UPI trades paid within 24 h"],
        ["Unevidenced rejections", "Arbitrary cuts are a major leak in income", "Zero rejections accepted without reason code and photo"],
        ["Unscheduled mandi deductions", "Opaque deductions erode net price", "Zero (every deduction maps to a published schedule line)"],
        ["Time to first lot (assisted)", "Onboarding friction kills adoption", "Under 5 minutes with an agent, under 10 minutes self-service"],
        ["Dispute resolution median", "Unresolved disputes freeze money", "Under 72 h from case open"],
        ["Farmers with a credit passport", "Credit exclusion is a root problem", "Share of active farmers who consent to export a verified history"],
    ], [1.7, 2.7, 2.45])

    # worked example
    h2(doc, "One trade, end to end (the example the platform must be able to run)")
    table(doc, ["#", "Moment", "What happens", "Rail"], [
        ["1", "Harvest morning", "Ramesh (farmer) declares 400 kg tomatoes from Plot 2 by app or voice. Platform creates the lot and QR and shows today's mandi band for his district.", "Goods"],
        ["2", "+10 min", "He opens Sell, picks Direct and scans retailer Sunita's shop QR. Sunita is verified with Pay-Score 86 (Reliable). Default terms: UPI on delivery. He sets Rs 20/kg: expected Rs 8,000.", "Money + Trust"],
        ["3", "+15 min (no network)", "His phone signs a Handover Token naming Sunita. QR appears on his screen. Nothing needs the internet.", "Goods"],
        ["4", "At the shop", "Sunita scans the token, weighs: 392 kg accepted, 6 kg rotten (reason + photo), 2 kg unaccounted. She countersigns on her phone.", "Goods + Trust"],
        ["5", "Network returns", "Both phones sync. Platform verifies two signatures and the one-time nonce. Tolerance check passes (2 kg is below the allowed tolerance). Events are written to the ledger.", "Ledger"],
        ["6", "Seconds later", "Split-pay: UPI request for 392 x Rs 20 = Rs 7,840 goes to Sunita. The 6 kg (Rs 120) is offered back to Ramesh: accept the rejection or contest within 24 h.", "Money"],
        ["7", "Minutes later", "Sunita pays; UTR is recorded. Ramesh taps Confirm on 'Rs 7,840 received'. Title moves to Sunita. Trade is SETTLED.", "Money"],
        ["8", "Same day", "Ramesh's delivery accuracy and Sunita's on-time payment update. A verified entry is added to his credit passport. Sunita prints consumer QR labels.", "Trust + Insight"],
        ["9", "Next day", "A shopper scans the pack: district, harvest date, two handovers, evidence level per claim. A freshness rating flows back to Ramesh as aggregated feedback.", "Insight"],
    ], [0.3, 1.1, 4.7, 0.75], size=8.5)

    # ---------------- 2 problems
    section(doc, landscape=True, margin=0.6)
    h1(doc, "Real farmer problems and the design response", "2")
    para(doc, "Each feature in this proposal exists to remove a specific, observable pain. Features that do not map to a problem here were left out.", size=9.5)
    table(doc, ["#", "Farmer problem", "Root cause in a traceability-only design", "Design response", "KPI"], [
        ["1", "Late, partial or unpaid sales", "Payment is outside the record; the buyer controls timing", "Trade object with payment terms, risk-based credit rule, split-pay, overdue ladder, Pay-Score shown before dispatch", "Hours from delivery to payment"],
        ["2", "Arbitrary rejection and weight cuts", "Only the buyer records weight and quality", "Dual-signed weigh slip, tolerance, mandatory reason + photo, rejection-anomaly detection, rejected qty stays with farmer", "Rejection rate vs peers"],
        ["3", "Opaque mandi deductions", "Commission-agent statements are unaudited", "Itemised account-sale; only scheduled deductions allowed; farmer sees it before payout", "Unscheduled deductions = 0"],
        ["4", "Price blindness", "Farmer has no reference at the moment of sale", "Mandi price band at listing, floor nudge, farmer-share view", "Price vs mandi modal"],
        ["5", "Closed networks and forced pre-linking", "Both sides must be registered and linked in advance", "Open handshake with provisional T0 accounts; trust earned, not pre-arranged", "Share of trades with new counterparties"],
        ["6", "Poor connectivity, low digital literacy", "App-only forms, always-online assumption", "Offline signed tokens, voice / IVR / WhatsApp, assisted entry by agents with consent", "Time to first lot"],
        ["7", "Credit exclusion", "No verified transaction history", "Consent-based credit passport via Account Aggregator; pay-today-against-receivable option", "Loans enabled"],
        ["8", "Spoilage and transit loss blamed on the farmer", "Custody and ownership conflated; no transit evidence", "Title vs custody split, carrier telemetry, loss attribution by delivery terms", "Transit loss %"],
        ["9", "No reward for quality or traceability", "Premium has no channel to flow back", "Verified-origin premium line item, quality score, consumer feedback loop", "Premium per kg"],
        ["10", "Data exploitation", "Everything visible to everyone in the chain", "DPDP-style consent, price visible only to the two trading parties, k-anonymous insights", "Consent coverage"],
    ], [0.3, 1.75, 2.35, 4.0, 1.7], size=8.5)

    # ---------------- 3 laws
    section(doc, landscape=False)
    h1(doc, "Twelve design laws", "3")
    para(doc, "These are the rules every later decision is checked against. When two laws conflict, the lower number wins.", size=9.5)
    laws = [
        ("The farmer is paid before anyone gets a dashboard.", "Money rail ships first."),
        ("A trade is a contract, not a status change.", "Price, quantity, terms and delivery are one signed object."),
        ("Title and custody are different things.", "The party holding goods is not necessarily the owner."),
        ("Pay the undisputed amount immediately.", "Only the contested part waits for resolution."),
        ("Both sides weigh; evidence beats assertion.", "No rejection without reason code and photo."),
        ("Offline is the normal case.", "Every field action must complete with no network and sync later."),
        ("No pre-linking.", "Any verified or provisional party can trade; trust is earned and visible."),
        ("A score is a ticket, never a gate.", "Scores unlock benefits; they never block a farmer from selling."),
        ("Data belongs to the farmer.", "Consent, minimum visibility, and the right to export."),
        ("The platform never holds the farmer's money.", "Payments move peer to peer through licensed rails."),
        ("Every claim shows its evidence level.", "Self-declared, peer-vouched or verified is always visible."),
        ("Boring technology, loud guarantees.", "Proven components; strong, testable invariants."),
    ]
    table(doc, ["#", "Law", "Meaning in practice"], [[str(i + 1), a, b] for i, (a, b) in enumerate(laws)],
          [0.35, 3.2, 3.3], size=9)

    # ---------------- 4 core concept
    h1(doc, "Core concept: four rails on one ledger", "4")
    para(doc, "Goods move, money settles, trust is earned and insight flows back. Each rail has its own objects and rules, but all four write to the same append-only, hash-chained ledger, "
              "so any dispute can be replayed from evidence.", size=9.5)
    table(doc, ["Object", "Question it answers", "Key fields", "Status in current build"], [
        ["Lot", "What physical batch?", "crop, qty, grade, plot, harvest date, origin claims with evidence level, owner, custodian", "Exists (add owner separate from custodian)"],
        ["Trade", "On what terms?", "seller, buyer, qty, price/kg, premium, payment terms, delivery terms, state", "NEW"],
        ["Custody", "Who physically holds it?", "custodian, pickup and drop events, telemetry", "Exists (handover)"],
        ["Title", "Who owns it?", "owner, transfer trigger (payment or acceptance)", "NEW"],
        ["Weigh slip", "What was counted, by whom?", "declared, accepted, rejected lines (reason, photo), both signatures", "NEW (partly in handover payload)"],
        ["Payment", "Was the farmer paid?", "amount, method (UPI, cash, bank), UTR, proof, confirmation by seller", "NEW"],
        ["Case", "What is contested?", "level, SLA, evidence pack, outcome", "Evolves from dispute"],
        ["Score", "Who is reliable?", "trust band, Pay-Score, quality score, evidence of components", "NEW"],
    ], [0.95, 1.5, 3.1, 1.3], size=8.5)
    callout(doc, "Why title and custody must be separate",
            "Today the system has one field, custodian. When a carrier holds goods, or a commission agent sells on consignment, that field cannot say whose goods they are, "
            "who bears loss, or who must be paid. Splitting title from custody lets a carrier hold goods without owning them, lets a farmer keep title until paid "
            "(for new buyers), and makes every loss attributable.", fill="F0FDF4")

    section(doc, landscape=True, margin=0.55)
    figure(doc, "01_four_rails.png", "Four rails: goods, money, trust and insight, tied together by dashed couplings", LW * 0.92)
    section(doc, landscape=True, margin=0.55)
    figure(doc, "02_system_layers.png", "System layers: clients, edge, domain services, data stores and integrations", LW * 0.92)

    # ---------------- 5 trade lifecycle
    section(doc, landscape=True, margin=0.55)
    h1(doc, "The Trade lifecycle", "5")
    figure(doc, "03_trade_states.png", "Trade state machine (numbers refer to the table on the next page)", LW)

    section(doc, landscape=False)
    h2(doc, "Transition table")
    table(doc, ["#", "From", "To", "Trigger", "Written to ledger"], [
        ["1", "DRAFT", "OFFERED", "Seller publishes price and terms", "TRADE_OFFERED"],
        ["2", "OFFERED", "AGREED", "Buyer accepts (or scans token)", "TRADE_AGREED"],
        ["3", "AGREED", "DISPATCHED", "Handover Token signed; pickup", "HANDOVER_OPEN"],
        ["4", "DISPATCHED", "DELIVERED", "Dual-signed weigh slip synced", "HANDOVER_RECEIVED, WEIGH_SLIP"],
        ["5", "DELIVERED", "PAYMENT_DUE", "Settlement computed (accepted qty x price + premium)", "SETTLEMENT"],
        ["6", "PAYMENT_DUE", "SETTLED", "Payment proof + seller confirmation (or PSP callback)", "PAYMENT_CONFIRMED, TITLE_TRANSFER"],
        ["7", "OFFERED", "CANCELLED", "Either party withdraws before agreement", "TRADE_CANCELLED"],
        ["8", "AGREED", "EXPIRED", "No dispatch before deadline", "TRADE_EXPIRED"],
        ["9", "DELIVERED", "DISPUTED", "Evidenced rejection or discrepancy beyond tolerance", "CASE_OPENED"],
        ["10", "DISPUTED", "PAYMENT_DUE", "Case resolved; adjusted amount agreed", "CASE_RESOLVED"],
        ["11", "PAYMENT_DUE", "OVERDUE", "Due date passed unpaid", "PAYMENT_OVERDUE"],
        ["12", "OVERDUE", "SETTLED", "Paid late; Pay-Score penalty applied", "PAYMENT_CONFIRMED_LATE"],
    ], [0.3, 1.15, 1.15, 2.6, 1.65], size=8.5)

    h2(doc, "Risk-based payment terms (the rule that protects the farmer)")
    para(doc, "The platform proposes the loosest terms that the buyer's standing justifies. The seller may always choose stricter terms. "
              "If the seller chooses looser terms than proposed, the app shows a clear warning and records 'seller accepted risk'. Thresholds below are starting values to tune in the pilot.", size=9.5)
    table(doc, ["Buyer standing", "Allowed terms", "Title moves when", "Reason"], [
        ["New, provisional (T0), or first trade between this pair", "Cash or UPI on delivery only", "Undisputed amount is paid", "No history; remove credit risk entirely"],
        ["Verified, Pay-Score below 80 or fewer than 5 trades with this seller", "UPI on delivery, or T+1", "Weigh slip accepted and payment claim recorded", "Some history; short exposure"],
        ["Proven: Pay-Score 80 or above, 5 or more trades with this seller", "T+3 up to credit limit", "Weigh slip accepted; receivable tracked", "Reliable record earns credit"],
        ["Institution or FPO-guaranteed buyer", "Per written agreement", "Per agreement; receivable tracked", "Contractual and reputational backing"],
        ["Any buyer with an overdue above 7 days", "Cash on delivery only", "Undisputed amount is paid", "Active default risk"],
    ], [2.0, 1.5, 1.8, 1.55], size=8.5)

    h2(doc, "Delivery terms (who bears transit risk)")
    table(doc, ["Term", "Risk passes to buyer", "Typical use"], [
        ["Ex-farm", "At pickup (carrier arranged by buyer)", "Retailer or wholesaler collects from the field"],
        ["Delivered", "At receiver's acceptance of the weigh slip", "Farmer or aggregator arranges transport"],
        ["Carrier-liable", "Per carrier agreement; loss beyond tolerance is claimed from the carrier", "Third-party logistics with telemetry"],
    ], [1.4, 2.8, 2.65], size=8.5)

    # ---------------- 6 handshake + settlement
    section(doc, landscape=False)
    h1(doc, "The universal handshake and settlement algorithm", "6")
    para(doc, "Every entity hands goods to another using the same protocol. It replaces the current requirement that both organisations be pre-linked as trusted partners. "
              "Instead, the two phones prove physical presence by exchanging signatures.", size=9.5)
    h2(doc, "Handover Token and receipt (offline-safe)")
    code(doc, """
HandoverToken (signed by seller device key)
  token_id, trade_id, lot_code
  declared_qty_kg, price_per_kg, premium_per_kg
  payment_terms, delivery_terms
  seller_id, buyer_id (named when seller scans buyer ID QR)
  nonce (single use), issued_at, expires_at
      default 72 h; 12-24 h for leafy and highly perishable crops

ReceiptToken (signed by buyer device key)
  token_id, accepted_qty_kg
  rejects[ {qty_kg, reason_code, photo_hash} ]
  condition, crates_counted, geo, time

Sync rules at the server
  1. Verify both signatures and that the buyer key matches buyer_id.
  2. Reject if nonce already consumed (idempotent by event_id).
  3. Apply events in custody-chain order (single-writer: only the
     current custodian may emit events for a lot).
  4. If either side is still offline, state stays DISPATCHED; goods
     and money are unaffected, only the record lags.
""", size=8)
    h2(doc, "Settlement algorithm")
    code(doc, """
function COMPUTE_SETTLEMENT(trade, slip):
    D = trade.declared_qty
    C = slip.accepted_qty
    R = sum(line.qty for line in slip.rejects)
    U = D - C - R                          # unaccounted quantity
    tol = max(2 * scale_resolution, k_crop * D)

    for line in slip.rejects:
        if not line.reason_code or not line.photo_hash:
            line.status = CLAIM_PENDING_EVIDENCE   # buyer has EVIDENCE_WINDOW
            C = C + line.qty                       # value stays payable now

    undisputed_value = C * (trade.price + trade.premium)
    pay_now(trade, undisputed_value)               # split-pay, never held

    if abs(U) > tol:
        open_case(trade, QTY_DISCREPANCY, U)       # tolerance exceeded
    for line in evidenced_rejects:
        offer_back_to_seller(line)                 # accept | reprice | take back | contest

    trade.due_date = delivered_at + trade.terms.days
    return Settlement(pay_now = undisputed_value, due = trade.due_date)
""", size=8)
    callout(doc, "Design note",
            "Tolerance still decides whether an alert or case opens, as in the current build. What changes is that tolerance no longer delays money: "
            "the undisputed value is always payable immediately, and the contested part is isolated.", fill="F0FDF4")
    section(doc, landscape=True, margin=0.55)
    figure(doc, "04_handshake.png", "Universal handshake: used by every entity, offline-safe, two signatures", LW)

    # ---------------- 7 entity chapters
    section(doc, landscape=False)
    h1(doc, "Entity-by-entity design", "7")
    para(doc, "Each entity has a promise, a boundary (what it does and never does), a pseudocode algorithm, enforced invariants, KPIs and a swimlane block diagram. "
              "All diagrams use one legend: green = goods or custody event, saffron = money event, blue = trust or evidence check, grey = ledger or system write, purple diamond = decision.",
         size=9.5)
    table(doc, ["#", "Entity", "One-line job", "Maps to current role"], [
        ["7.1", "Producer (farmer)", "Create honest lots and get paid fast", "FARMER / Producer"],
        ["7.2", "Aggregator / FPO", "Pool lots fairly, pay same day, share upside", "AGGREGATOR"],
        ["7.3", "Wholesaler / commission agent", "Run an auditable market and statement", "WHOLESALER"],
        ["7.4", "Distributor / carrier", "Hold custody (not title), prove condition", "DISTRIBUTOR / OPERATOR"],
        ["7.5", "Retailer", "Buy from anyone verified, sell with consumer QR", "RETAILER"],
        ["7.6", "Consumer", "See provenance, give feedback, thank the grower", "Public view"],
        ["7.7", "Verifier / mediator / admin", "Keep the network honest and fast", "Admin, LAB"],
    ], [0.45, 2.0, 2.9, 1.5], size=9)

    entity_chapter(
        doc, "7.1", "Producer (farmer)",
        "Sell at a fair price, get paid fast, and keep my proof. I should never need a computer, a pre-arranged buyer, or a lawyer.",
        ["Register by phone number (or via an agent with consent)", "Declare harvest and create lots with QR",
         "See the mandi price band, set price and terms", "Hand over by scan and offline token",
         "Confirm payment received (tap, OTP or voice)", "Export a consent-based credit passport"],
        ["Pre-link buyers before selling", "Pay any fee for the core flow", "Accept a rejection that lacks reason and photo",
         "Have plot geometry shown publicly", "Have the price overwritten by anyone downstream"],
        ["Farmer-set price is never changed by a downstream party.",
         "Rejected quantity stays under the farmer's title until a disposal is agreed.",
         "A trade is SETTLED only after farmer confirmation or a PSP-verified payment.",
         "An agent-assisted entry needs farmer consent (OTP, voice or thumb) within 24 h or is flagged.",
         "Yield plausibility: qty per plot per season cannot exceed area x crop max yield x slack."],
        ["Time to first lot (assisted and self-service)", "Median hours from delivery to confirmed payment",
         "Share of trades settled without dispute", "Realised price vs mandi modal price"],
        """
procedure SELL_LOT(farmer, lot):
    ref   = MANDI_REFERENCE(lot.crop, lot.district)       # Agmarknet / eNAM band
    route = farmer.choose(DIRECT | FPO | MANDI)
    buyer = farmer.pick_or_scan(route)                    # no pre-link needed
    assert buyer.tier >= required_tier(lot.qty)           # provisional T0 allowed
    terms = DEFAULT_TERMS(farmer, buyer)                  # risk-based (Section 5)
    price = farmer.set_price(prefill = ref.modal)
    warn_if(price < 0.9 * ref.min)                        # floor nudge, not a block
    trade = CREATE_TRADE(lot, buyer, qty, price, terms, delivery_terms)
    token = SIGN_HANDOVER_TOKEN(trade, buyer.pubkey, ttl) # works offline
    on receipt_token(slip):
        s = COMPUTE_SETTLEMENT(trade, slip)               # split-pay now
        notify(farmer, s)
    on payment_marked(utr):
        farmer.confirm() or PSP_VERIFY(utr)
    UPDATE_SCORES(trade); ADD_TO_CREDIT_PASSPORT(trade)
""",
        "05_producer.png", "Producer algorithm swimlane: from harvest to paid-and-confirmed")

    entity_chapter(
        doc, "7.2", "Aggregator / FPO collection centre",
        "Let me pool many small farmers into sellable lots, pay them the same day, keep a transparent margin, and share upside fairly.",
        ["Publish a daily rate card by grade before opening", "Weigh in front of the farmer with a digital slip and photo",
         "Pay farmers same day (UPI batch or cash with OTP receipt)", "Merge lots while preserving each farmer's contribution share",
         "Sort into grade lots with new QR labels", "Sell through the standard sell flow"],
        ["Change the rate card mid-day without publishing", "Pay below the card without evidence",
         "Hide who contributed what to a pooled lot", "Hold farmer money for later sale outcomes"],
        ["For every pooled lot, the sum of farmer shares equals 1 and each farmer's slip is retained.",
         "A pay-SLA breach by the aggregator triggers the same overdue ladder as any buyer.",
         "Margin cap per FPO charter is configurable; excess over rate paid plus cap is shared pro-rata.",
         "Re-grading needs evidence; the farmer may accept, withdraw or ask mediation."],
        ["Share of farmers paid same day", "Collection-to-sale time", "Premium share distributed per season", "Re-grade dispute rate"],
        """
procedure COLLECTION_DAY(centre):
    publish(RATE_CARD(crop, grade, today))                # before opening
    for each farmer arrival:
        slip  = WEIGH_BOTH_SIDES(farmer.lot)              # photo, time, geo
        grade = QUALITY_CHECK(slip)                       # moisture, size, defects
        if grade.meets(rate_card):
            rate    = rate_card[grade]
            payable = slip.net_qty * rate
            TRADE(farmer -> centre, payable, terms = SAME_DAY)
            PAY_NOW(farmer, payable)                      # batch UPI or cash + OTP
        else:
            REGRADE_WITH_EVIDENCE(farmer, slip)           # accept | withdraw | mediate
    at cutoff:
        pool = MERGE(lots, keep_shares = TRUE)            # farmer share ledger
        for lot in SORT_GRADE(pool): SELL_LOT(centre, lot)
    for lot in sold_lots:
        excess = lot.realised - lot.rate_paid - centre.margin_cap
        if excess > 0: SHARE(excess, farmers, by = contribution_share)
""",
        "06_aggregator.png", "Aggregator / FPO algorithm swimlane: same-day pay, fair pooling, shared upside")

    entity_chapter(
        doc, "7.3", "Wholesaler / commission agent",
        "Let me run a fast market where every bid, every deduction and every payout is on the record, so my reputation becomes an asset instead of a rumour.",
        ["Receive consignments with weigh, temperature and quality checks", "Operate in principal or commission mode",
         "Log every auction bid with time", "Split into sale lots and issue new QR labels",
         "Issue an itemised account-sale statement before payout", "Pay net to the seller within the published SLA"],
        ["Deduct any charge not on the published schedule", "Pay out before the seller has seen the statement",
         "Sell on commission without a logged bid", "Change the schedule mid-season without notice"],
        ["gross = sum of sale legs; net = gross minus the sum of scheduled deductions.",
         "Each deduction line carries a schedule ID and a receipt; unscheduled lines are blocked.",
         "The seller sees the statement before payout and may query any line.",
         "Payout SLA per mandi (T+X) is configured; breach opens the overdue ladder."],
        ["Unscheduled deductions (target zero)", "Hours from sale to account-sale statement", "Dispute rate per mandi", "Median payout delay"],
        """
procedure MANDI_SALE(agent, lot):
    receive(lot) via HANDSHAKE                            # weigh, temp, quality
    mode = agent.mode(lot)                                # PRINCIPAL | COMMISSION
    if mode == COMMISSION:
        log = AUCTION_LOG(lot)                            # bidder, price, time
        sale_price = log.winning_bid
        for split in SPLIT(lot, buyers):
            TRADE(agent -> buyer, split)                  # one trade per split
        stmt = ACCOUNT_SALE(
            gross      = sum(split.value),
            deductions = [d for d in charges if d in SCHEDULE(mandi)],
            blocked    = [d for d in charges if d not in SCHEDULE(mandi)])
        assert stmt.blocked == []
        SHOW_TO_SELLER(stmt); WAIT(seller.approve_or_query)
        PAY(seller, stmt.net, within = mandi.payout_sla)
    else:                                                 # PRINCIPAL
        TRADE(seller -> agent, price)                     # agent buys, takes title
        PAY(seller, per_terms)
        resell via SELL_LOT
    UPDATE_SCORES(); DEDUCTION_ANALYTICS(mandi)
""",
        "07_wholesaler.png", "Wholesaler / commission agent algorithm swimlane: auditable statement before payout")

    entity_chapter(
        doc, "7.4", "Distributor / carrier",
        "Let me move produce safely, prove its condition at every kilometre, and be judged on what actually happened to it in my custody.",
        ["Count crates, photograph and read temperature at pickup", "Log GPS, temperature and door-open events (offline buffer)",
         "Hand over via the same handshake at drop", "Attach a freight invoice to the trade it served"],
        ["Take title to the goods", "Settle transit loss with the farmer", "Edit telemetry after the fact",
         "Mark a leg closed without receiver counts"],
        ["Custody is not title: a carrier holds goods without owning them.",
         "Every leg has pickup and drop evidence (counts, weight where available, temperature).",
         "Cold-chain or route breach raises an alert to owner and receiver immediately.",
         "Transit loss beyond tolerance is attributed by delivery terms and claimed from the carrier when the carrier was in custody."],
        ["Transit loss percent per route", "Cold-chain breach rate", "On-time delivery", "Claims resolved within SLA"],
        """
procedure CARRY(carrier, leg):
    PICKUP: count crates, photo, temperature -> PICKUP_EVENT   # custody only
    while in_transit:
        every N minutes: LOG(gps, temp, door_events)           # buffer offline
        if temp > limit for > t or route_dev > x or delay > ETA + slack:
            ALERT(owner, receiver); ADJUST_SHELF_LIFE(lot)
    DELIVER: receiver HANDSHAKE (recount, weigh)
    loss = pickup_qty - drop_qty
    if abs(loss) <= tolerance(leg):
        CLOSE_LEG(leg); LINK_FREIGHT_INVOICE(leg.trade)
    else:
        attribute = BY_DELIVERY_TERMS(leg)                     # ex-farm | delivered | carrier-liable
        OPEN_CLAIM(attribute.party, loss)                      # never silently the farmer
""",
        "08_carrier.png", "Distributor / carrier algorithm swimlane: custody without title")

    entity_chapter(
        doc, "7.5", "Retailer",
        "Let me buy fresh produce straight from a farmer or any verified supplier, sell it with a consumer QR, and show I treat suppliers fairly.",
        ["Receive with weigh, temperature, quality and shelf-life check", "Accept or reject with reason code and photo",
         "Pay per terms and show the SLA clock to the supplier", "Split to shelf batches and print consumer QR labels",
         "Log waste and demand signals", "Respond to recalls"],
        ["Reject without reason and photo", "Hold rejected goods without an agreed disposal",
         "Ignore a recall notice", "Buy only from pre-linked suppliers"],
        ["Rejected quantity stays under the supplier's title until disposal is agreed.",
         "Payment due date is visible to the supplier from the moment of delivery.",
         "A recall notice must be acknowledged within SLA; unacknowledged notices escalate.",
         "Rejection rate far above peer median for the same crop triggers an anomaly flag and lowers Pay-Score."],
        ["Rejection rate vs peers", "Days payable", "Waste percent", "Recall response time"],
        """
procedure RETAIL(retailer):
    scan supplier QR; SIGN receipt via HANDSHAKE          # direct from farmer OK
    check = weigh + temperature + quality + shelf_life_remaining(lot)
    if check.fails(min_shelf_life):
        REJECT(lot, reason_code, photo)                   # stays with supplier
    else:
        PAY(supplier, per_terms)                          # UPI or cash + OTP
    split lot into shelf batches; PRINT consumer QR with best_before
    sell at POS:
        decrement on_hand
        if shelf_life < threshold: MARK_DOWN(batch)
        log waste(reason) -> demand_signal(supplier)
    on RECALL(lot):
        QUARANTINE(batches); BLOCK_POS(batches); NOTIFY(upstream)
    if lot.verified_origin and premium_agreed:
        premium_line_item credited to supplier in the trade
""",
        "09_retailer.png", "Retailer algorithm swimlane: fair sourcing, labelling and recalls")

    entity_chapter(
        doc, "7.6", "Consumer",
        "Let me know what I am eating and who grew it, with nothing hidden and nothing about the farmer that they did not agree to share.",
        ["Scan the QR on a pack", "See origin district, harvest date, journey and evidence level per claim",
         "See an anonymised farmer-share indicator", "Rate freshness, report an issue, optionally send a UPI thank-you"],
        ["See prices of any trade", "See farmer phone, ID or plot geometry",
         "Trigger a recall alone (reports are clustered and reviewed)"],
        ["Public view is built only from consented, zero-PII read models.",
         "Farmer-share indicator shows only when enough trades exist (k-anonymity).",
         "Issue reports are rate-limited and clustered by lot to resist abuse.",
         "Thank-you tips go UPI direct to the farmer with the farmer's consent."],
        ["Scans per lot", "Report-to-review time", "Share of scanners who rate or tip", "Feedback delivered to farmers per week"],
        """
procedure CONSUMER_SCAN(qr):
    lot  = RESOLVE(qr)
    view = PUBLIC_VIEW(lot)       # consented fields only; no prices; no PII
    show(view.origin_district, view.harvest_date, view.journey)
    for claim in view.claims: show(claim.text, claim.evidence_level)
    if k_anonymity_ok(lot): show(FARMER_SHARE(lot))       # aggregated band
    actions = [RATE_FRESHNESS, THANK_FARMER_UPI, REPORT_ISSUE]
    if action == REPORT_ISSUE:
        signal = CLUSTER(report, lot)
        if signal.count >= threshold or report.severity == SAFETY:
            QUEUE_FOR_VERIFIER(lot)                       # may open recall
    AGGREGATE_FEEDBACK(lot) -> farmer, retailer           # anonymous, weekly
""",
        "10_consumer.png", "Consumer algorithm swimlane: provenance, feedback and signals")

    entity_chapter(
        doc, "7.7", "Verifier / mediator / admin (and lab)",
        "Let me keep the network honest and fast: onboard in minutes, catch fraud early, resolve disputes before they freeze money.",
        ["Review flagged registrations and vouch for new farmers", "Sample re-verification of plots and facilities",
         "Mediate cases using an auto-built evidence pack", "Open and track recalls",
         "Attach lab results as high-evidence claims"],
        ["Reveal personal data without a logged reason", "Edit or delete a ledger event",
         "Decide a case without reading the evidence pack", "Approve their own relatives"],
        ["Ledger is append-only; corrections are new events that supersede, never edits.",
         "Every sensitive reveal is audit-logged (already built).",
         "Tier-3 approvals need dual control.",
         "Verification SLA by tier; breaches are visible to the applicant."],
        ["Verification SLA met", "Dispute resolution median", "Fraud-flag precision (true flags / all flags)", "Recall completion time"],
        """
procedure ONBOARD(applicant):
    flags = AUTO_CHECKS(applicant)
        # duplicate ID or plot, geo mismatch, qty > area * max_yield * slack,
        # shared device or phone clusters, velocity spike vs 90-day baseline
    if flags.empty: APPROVE_TIER(applicant, tier = T0_or_T1); ISSUE_QR_AND_CODE()
    else:           QUEUE_HUMAN_REVIEW(applicant, sla = SLA[tier])
    decision = verifier.decide(MORE_INFO | APPROVE | REJECT)
    notify(applicant, reason_in_local_language)

procedure MEDIATE(case):
    pack = BUILD_EVIDENCE_PACK(case)     # slips, photos, geo, timeline, scores
    proposal = RULES_PROPOSAL(pack)      # then human adjusts within policy
    if parties.accept(proposal): CLOSE(case, outcome)
    else:                         ESCALATE(case, APMC_or_ODR)
    WRITE_LEDGER(CASE_RESOLVED); UPDATE_SCORES(case.parties)
""",
        "11_verifier.png", "Verifier / mediator onboarding algorithm: progressive trust, fraud rules before humans")

    # ---------------- 8 special flows
    section(doc, landscape=False)
    h1(doc, "Special flows", "8")
    h2(doc, "8.1 Farmer sells directly to a retailer")
    para(doc, "This is the case the current build makes awkward: the code does not check roles, but it requires a bilateral trusted-partner link, so a farmer and a shop that have never met cannot trade without a registration and invitation dance. The proposed flow removes that step.", size=9.5)
    bullets(doc, [
        "Scan, do not search:: the farmer scans the retailer's entity QR (shop sign or phone). The retailer need not be pre-registered.",
        "Provisional account:: an unknown retailer gets a provisional T0 account by OTP in about 30 seconds, limited to cash or UPI on delivery.",
        "Defaults protect the farmer:: terms are prefilled from the mandi band and the risk rule; credit is unavailable for provisional buyers.",
        "Same-day money:: split-pay requests the full undisputed value at weighing; only a rejected part can be contested.",
        "The record is the same:: a two-step journey on the consumer page, with the same ledger guarantees.",
    ], size=9.5)
    section(doc, landscape=True, margin=0.55)
    figure(doc, "12_direct_sale.png", "Direct farmer-to-retailer sale: no pre-linking, same-day money", LW)
    section(doc, landscape=True, margin=0.55)
    figure(doc, "13_overdue_ladder.png", "Payment overdue escalation ladder", LW)
    section(doc, landscape=False)
    h2(doc, "8.2 Overdue ladder rules")
    para(doc, "The ladder applies only to credit and delayed terms. Dates are relative to the agreed due date D. All timings are starting values.", size=9.5)
    table(doc, ["When", "Automatic action", "Who is told"], [
        ["D-1", "Reminder with one-tap pay link", "Buyer"],
        ["D+0", "Mark OVERDUE; show on both parties' screens", "Buyer, farmer"],
        ["D+2", "Open a mediation case automatically", "Mediator, both parties"],
        ["D+7", "Buyer limited to cash terms; warning shown to farmers viewing this buyer", "Buyer, prospective sellers"],
        ["D+15", "Generate formal notice pack: ledger extract and acknowledgement of debt", "Farmer, buyer"],
        ["D+30", "Network default flag (buyer may contest with evidence)", "Network (band only)"],
    ], [0.8, 4.2, 1.85], size=9)
    callout(doc, "Legal caution",
            "The legal weight of an acknowledgement of debt or a notice pack depends on jurisdiction and drafting. Have counsel review the templates before relying on them.",
            fill="FFF7E6", bar="C2670A")
    section(doc, landscape=True, margin=0.55)
    figure(doc, "14_dispute_ladder.png", "Dispute resolution ladder: cheapest, fastest fix first", LW)

    # ---------------- 9 trust
    section(doc, landscape=False)
    h1(doc, "Trust and reputation algorithms", "9")
    callout(doc, "Principle", "A score is a ticket, never a gate. It can unlock credit, premiums and faster terms. It must never stop a verified farmer from selling or hide them from buyers.")
    h2(doc, "Producer trust score (0 to 100)")
    code(doc, """
T = 100 * ( 0.30*D + 0.25*Q + 0.20*V + 0.15*R + 0.10*F )

D = delivery accuracy   share of trades delivered within tolerance and on time
Q = quality consistency 1 - rejection rate of this farmer's lots, normalised
                        against peers for the same crop and season
V = verification depth  T0 = 0.25, T1 = 0.50, T2 = 0.80, T3 = 1.00
R = reliability         share of committed volume actually delivered
F = fairness            disputes settled amicably or upheld, vs claims found false

History weighting: exponential, 90-day half-life.
Shown only after 5 trades; before that the band is NEW.
""", size=8)
    h2(doc, "Buyer Pay-Score (0 to 100), shown to farmers before dispatch")
    code(doc, """
P = 100 * ( 0.45*O + 0.20*U + 0.15*L + 0.10*N + 0.10*V )

O = on-time payment rate (by value)
U = undisputed rate      1 - abusive-rejection ratio (rejection rate far above
                         peer median for the same crop is penalised)
L = lateness severity    1 - min(1, mean_days_late / 7)
N = history depth        min(1, completed_trades / 30)
V = verification depth   as above

credit_limit = base_limit(tier) * (P / 100)^2        # only if P >= 60 and no overdue > 7 days
First trade between any pair is always cash or UPI on delivery.
""", size=8)
    table(doc, ["Band", "Score", "Shown to the other side as"], [
        ["Reliable", "80 and above", "Green badge plus on-time percent"],
        ["Good", "60 to 79", "Blue badge plus on-time percent"],
        ["Watch", "40 to 59", "Amber badge plus on-time percent and last overdue"],
        ["Risky", "Below 40", "Red badge plus on-time percent and open overdues"],
        ["New", "Fewer than 5 trades", "Grey badge; cash or UPI on delivery only"],
    ], [1.3, 1.6, 3.95], size=9)
    h2(doc, "Fraud and plausibility rules (run before any human sees a case)")
    bullets(doc, [
        "Yield plausibility:: qty per plot per season must not exceed area x crop maximum yield x slack.",
        "Geo consistency:: lot GPS far from the plot centroid is flagged.",
        "Velocity:: volume above 3x the 90-day baseline goes to review.",
        "Circular trading:: repeated trades between the same pair with no onward movement, or parties sharing a device or phone cluster, are flagged.",
        "Price outliers:: prices far above or below the mandi band (either direction) are flagged for review, not blocked.",
        "Rejection anomaly:: a buyer's rejection rate above twice the peer median for the same crop in 30 days lowers U and alerts a mediator.",
    ], size=9.5)
    para(doc, "Fairness safeguards: every score shows its components, farmers can appeal a component, weather-driven variation is normalised against peers, and a yearly bias review checks outcomes across gender, land size and region.", size=9.5)

    # ---------------- 10 money data compliance
    h1(doc, "Money, data and compliance", "10")
    h2(doc, "Payment architecture (phased, never holding funds)")
    table(doc, ["Phase", "How money moves", "How the platform knows", "Limits"], [
        ["1: Record and nudge", "Platform creates a UPI intent link or QR (payee = seller, amount, reference = trade ID). Cash is also supported.", "Buyer enters UTR; seller confirms by tap, OTP or voice", "Relies on seller confirmation; fake UTR is caught by the seller"],
        ["2: Auto-reconcile", "Same, via a licensed PSP payment link or collect request", "PSP webhook marks paid in seconds; reconciliation report", "Needs PSP partner and merchant onboarding"],
        ["3: Assured settlement (optional)", "For credit sales, a regulated lender may pay the farmer today against the buyer's receivable", "Receivable and Pay-Score data, consent-based", "Platform does not lend; lender terms apply"],
    ], [1.2, 2.3, 1.8, 1.55], size=8.5)
    h2(doc, "Who can see what")
    table(doc, ["Data", "Farmer", "Direct counterparty", "Other chain members", "Consumer", "Verifier / mediator"], [
        ["Trade price", "Own trades", "Own trade only", "No", "No", "Only in an open case"],
        ["Quantities, weigh slips", "Own", "Own trade", "Own legs", "No", "Case only"],
        ["Farmer name and photo", "Self", "Yes", "First name and district", "First name and district, only with consent", "Yes (reveal logged)"],
        ["Plot geometry", "Self", "No", "No", "District or block only", "T2+ verification only"],
        ["Phone and ID documents", "Self", "Masked phone", "No", "No", "Reveal logged"],
        ["Scores", "Full components", "Band + on-time %", "Band", "No", "Full, in a case"],
    ], [1.25, 0.85, 1.05, 1.2, 1.3, 1.2], size=8)
    h2(doc, "Compliance checklist (to validate with counsel before launch)")
    table(doc, ["Area", "Design position", "Action needed"], [
        ["Personal data (DPDP Act 2023)", "Consent per purpose, minimum visibility, export and erasure for non-ledger data; ledger stores pseudonymous references", "Privacy notice in local languages; consent manager design"],
        ["Payments (RBI rules)", "Platform holds no funds; UPI and licensed PSPs only; any escrow via an authorised arrangement", "Legal review and PSP agreement"],
        ["Mandi rules (state APMC)", "Deduction schedules are configured per mandi and state", "Obtain schedules for the pilot mandi"],
        ["Food safety traceability (FSSAI)", "Batch codes, supplier records and retention aligned to current requirements", "Check retention periods and label content"],
        ["GST and e-way bill", "Link e-way bill numbers on movements; platform does not file returns", "Confirm integration scope"],
        ["Electronic records as evidence", "Signed, timestamped, hash-chained records with audit trail", "Counsel review of notice and acknowledgement templates"],
    ], [1.7, 3.2, 1.95], size=8.5)

    # ---------------- 11 data model
    section(doc, landscape=True, margin=0.55)
    h1(doc, "Data model: what is new and how it connects", "11")
    figure(doc, "15_data_model.png", "Data model: new objects wrap existing tables so money and evidence sit beside goods", LW * 0.85)
    section(doc, landscape=False)
    h2(doc, "Keep, evolve or add")
    table(doc, ["Current element", "Decision", "Change"], [
        ["Hash-chained trace events and daily Merkle anchor", "Keep", "Add event types: TRADE_*, PAYMENT_*, TITLE_TRANSFER, WEIGH_SLIP, REJECT, REPRICE, CASE_*"],
        ["Lot, lot edges (split, merge, sort), mass balance", "Keep + extend", "Add owner_org separate from custodian_org; origin claims with evidence level"],
        ["Handover (two-step custody)", "Evolve", "Becomes the delivery step of a Trade; carries token_id; no mandatory partner link"],
        ["PartnerLink (bilateral trusted partner)", "Evolve", "Optional 'preferred partner' relationship for credit terms and auto-accept; no longer required to trade"],
        ["Tolerance and discrepancy alerts", "Keep + extend", "Tolerance triggers alert or case but no longer delays undisputed payment"],
        ["Dispute states", "Evolve", "Four-level ladder with SLAs and evidence pack"],
        ["OTP and passkey identity, consent records, encrypted sensitive data", "Keep + extend", "Add tiers T0-T3, delegate grants, device keys for signing"],
        ["Offline sync batch endpoint", "Evolve", "Accepts signed tokens and receipts; idempotent by event_id; nonce single-use"],
        ["Public view read model, QR labels", "Keep + extend", "Add evidence level per claim, farmer-share indicator, feedback"],
        ["11-language UI", "Keep + extend", "Add voice and IVR prompts; same string keys"],
        ["Trade, TradeLine, Payment, WeighSlip, Evidence, Case, ScoreSnapshot, Delegate, DeviceKey, HandoverToken, AccountSale, PriceReference", "Add", "New tables; see figure above"],
    ], [2.6, 1.0, 3.25], size=8.5)

    # ---------------- 12 incentives
    h1(doc, "Incentives and business model (hypotheses)", "12")
    para(doc, "The farmer never pays for the core flow. Others pay because the platform lowers their cost or risk.", size=9.5)
    table(doc, ["Who pays", "For what", "Why they would"], [
        ["Retailers and brands", "Verified-origin premium; traceability and recall readiness; freshness and waste analytics", "Fresher produce, fewer recalls, marketing value"],
        ["Aggregators, wholesalers, agents", "Software for settlement, account-sale and score portability", "Fewer disputes, faster working-capital turns, access to credit"],
        ["Lenders and insurers", "Consent-based access to verified transaction history", "Lower underwriting cost and loss rates"],
        ["Government and FPO programmes", "Rollout, assisted-entry agents, mandi digitisation", "Measurable farmer income and loss reduction"],
        ["Consumers (optional)", "UPI thank-you to the grower", "Goodwill; fully optional"],
    ], [1.7, 2.9, 2.25], size=8.5)

    # ---------------- 13 roadmap and pilot
    section(doc, landscape=True, margin=0.55)
    h1(doc, "Roadmap, pilot and success measures", "13")
    figure(doc, "16_roadmap.png", "Four phases, each shippable on its own (durations are indicative)", LW * 0.92)
    section(doc, landscape=False)
    h2(doc, "Pilot design")
    table(doc, ["Item", "Proposal"], [
        ["Crop", "One perishable crop with many intermediaries and visible spoilage (for example tomato)"],
        ["Geography", "One district cluster"],
        ["Participants", "One FPO (100 to 300 farmers), 2 to 3 aggregators, 1 mandi with about 5 commission agents, 15 to 25 retailers, 2 carriers"],
        ["Duration", "2 weeks baseline measurement, then 10 to 12 weeks live"],
        ["Comparison", "Pilot farmers vs similar non-pilot neighbours on payment delay, rejection percent, deductions and net price"],
        ["Go / no-go", "Median delivery-to-payment under 24 h; unevidenced rejections zero; at least 70% of assisted farmers transact again within 30 days"],
    ], [1.3, 5.55], size=9, first_col_bold=True)
    h2(doc, "Phase exit criteria")
    table(doc, ["Phase", "Exit criteria"], [
        ["1: Money rail", "Every pilot trade has terms and a payment record; split-pay works; ledger invariants unchanged; no data loss"],
        ["2: Field rail", "Offline tokens sync with zero double-spend in tests; assisted entry live; account-sale statements issued; dispute ladder SLAs measured"],
        ["3: Value rail", "PSP reconciliation live; first credit-passport exports with consent; price board in use"],
    ], [1.3, 5.55], size=9, first_col_bold=True)

    # ---------------- 14 risks
    h1(doc, "Known weaknesses, risks and open decisions", "14")
    table(doc, ["Risk", "How it shows up", "Containment"], [
        ["Low farmer adoption", "Forms, language, no smartphone", "Assisted agents, voice and IVR, FPO champions, immediate money benefit"],
        ["Intermediary resistance", "Commission agents resist transparency", "Start with willing early adopters; portable Pay-Score gives them credit access; align with mandi authority"],
        ["Fake data and collusion", "False lots, circular trading", "Plausibility and anomaly rules, evidence levels, sample audits, device clustering"],
        ["Payment fraud", "Fake UTR, false 'paid' claims", "Seller confirmation; PSP webhook in Phase 2"],
        ["Regulatory exposure", "Holding funds, legal weight of notices", "No funds held; licensed partners; counsel review"],
        ["Connectivity", "No network at farm or mandi", "Offline tokens, SMS fallback, store-and-forward"],
        ["Score misuse or bias", "Excluding small or remote farmers", "Ticket-not-gate law, appeals, normalisation, yearly bias audit"],
        ["Privacy breach", "PII exposure", "Encryption, minimum visibility, reveal logging (existing), DPDP processes"],
        ["Complexity creep", "Too many features at once", "Four phases, feature flags, boring technology"],
        ["Third-party dependency", "Price feed or PSP outage", "Cache with last-known band; manual UTR fallback"],
    ], [1.5, 2.2, 3.15], size=8.5)
    h2(doc, "Open decisions for stakeholders")
    bullets(doc, [
        "Operator and liability:: who is the legal operator and who is liable for dispute outcomes?",
        "Pilot jurisdiction:: which state APMC rules and which mandi schedule apply?",
        "PSP partner:: which licensed payment provider and what onboarding timeline?",
        "Discovery:: adopt an open commerce protocol for buyer discovery, or stay with scan-to-trade first?",
        "Premium funding:: who funds the verified-origin premium and how is it priced?",
        "Mediator panel:: composition, training and neutrality safeguards.",
    ], size=9.5)
    h2(doc, "Assumptions to validate with farmers before building")
    bullets(doc, [
        "How long do farmers wait for payment today, by buyer type, and what do they lose when it is late?",
        "What share have a working UPI handle and use it?",
        "How is weighing done at the farm gate, collection centre and mandi, and who owns the scale?",
        "Would farmers rather confirm payment by tap, OTP or voice call?",
        "What are they willing to share (plot map, ID) and with whom?",
        "Does an FPO or agent they already trust exist locally to run assisted entry?",
    ], size=9.5, bold_lead=False)

    # ---------------- appendix
    h1(doc, "Appendix A: API sketch", "A")
    table(doc, ["Endpoint", "Purpose"], [
        ["POST /v1/trades", "Create an offer on a lot"],
        ["POST /v1/trades/{id}/accept", "Buyer accepts"],
        ["POST /v1/trades/{id}/token", "Issue a signed handover token (also works offline on device)"],
        ["POST /v1/sync/events", "Batch, idempotent upload of signed events and receipts"],
        ["POST /v1/trades/{id}/weigh-slip", "Submit or countersign a weigh slip"],
        ["POST /v1/trades/{id}/payments", "Record a payment (method, UTR, proof)"],
        ["POST /v1/payments/{id}/confirm", "Seller confirms receipt"],
        ["POST /v1/trades/{id}/cases", "Open a case with evidence"],
        ["GET /v1/price-reference?crop=&district=", "Mandi price band"],
        ["GET /v1/entities/{id}/score", "Band and components, visibility-filtered"],
        ["GET /v1/me/credit-passport", "Consent-gated verified history"],
        ["POST /v1/delegates", "Grant an agent authority to act for a farmer"],
    ], [3.0, 3.85], size=8.5)
    h1(doc, "Appendix B: Glossary", "B")
    table(doc, ["Term", "Meaning"], [
        ["Lot", "A physical batch of produce with a QR code"],
        ["Trade", "A contract between seller and buyer on a lot quantity with price and terms"],
        ["Custody", "Physical possession of goods"],
        ["Title", "Legal ownership of goods"],
        ["Handshake", "The two-signature, offline-safe handover protocol"],
        ["Token", "A signed, single-use record exchanged at handover"],
        ["Split-pay", "Paying the undisputed value immediately while isolating the contested part"],
        ["Pay-Score", "Buyer reliability score shown to farmers"],
        ["Tier (T0-T3)", "Verification depth: phone, identity or vouch, land or licence, audited"],
        ["Delegate", "An agent authorised by a farmer to act on their behalf"],
        ["Account-sale", "Itemised statement of a commission sale: gross minus scheduled deductions"],
        ["Credit passport", "Consent-based export of verified transaction history"],
    ], [1.6, 5.25], size=9, first_col_bold=True)

    os.makedirs(os.path.dirname(OUT_DOCS), exist_ok=True)
    doc.save(OUT_DOCS)
    shutil.copyfile(OUT_DOCS, OUT_ROOT)
    print("saved", OUT_DOCS)


if __name__ == "__main__":
    build()
